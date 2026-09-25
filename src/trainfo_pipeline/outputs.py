"""Write reproducible CSV and formatted Excel pipeline outputs."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_excel_table(frame: pd.DataFrame, path: Path, title: str, updated: str) -> None:
    """Write a titled Excel table with a visible last-updated date."""
    frame.to_excel(path, index=False, startrow=3, engine="openpyxl")
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.title = "Trainfo Data"
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"Last updated: {updated}"
    sheet["A2"].font = Font(italic=True)
    end_row = len(frame) + 4
    table = Table(displayName="TrainfoData", ref=f"A4:I{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[letter].width = min(max(width, 12), 42)
    workbook.save(path)


def write_dataset_outputs(
    frame: pd.DataFrame,
    output_dir: Path,
    label: str,
    updated: date,
) -> None:
    """Write latest and dated CSV/XLSX copies for one dataset."""
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = updated.isoformat()
    date_label = updated.strftime("%B %-d, %Y")
    title = f"Trainfo Houston - {label} - Last Updated {date_label}"
    for suffix in ("latest", stamp):
        frame.to_csv(output_dir / f"trainfo_houston_{label}_{suffix}.csv", index=False, encoding="utf-8-sig")
        write_excel_table(
            frame,
            output_dir / f"trainfo_houston_{label}_{suffix}.xlsx",
            title,
            date_label,
        )
